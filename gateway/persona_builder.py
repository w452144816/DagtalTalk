"""多模态素材 → 结构化人设（docs/AI服务多机方案.md 第十节 10.2），只在 control 角色挂载。

人设 `ps_xxx` 是独立对象：多份素材（文档 / 图片 / 音频 / 视频 / 网页）各自抽取成要点，`build` 合并成一个结构化 JSON
并编译出 system_prompt；每次 build / 人工改都 +1 版本，会话引用某个版本。draft 可文本试聊，approve 后才给正式会话用。

抽取用什么：文本类直接读（PDF→pypdf、Word→python-docx、网页→bs4）；图片 → 方舟多模态模型描述 + OCR；
音频 / 视频 → ffmpeg 转 16k PCM 分段送节点 SenseVoice 识别（视频再抽 3 帧看着装场景）；然后 LLM 按 schema 抽要点。
知识量：编译后 ≤ 3000 token 内联到 prompt（inline），超过只内联人设/边界，知识导出给后端做 RAG（external）。
落盘：/w/state/personas/<pid>/persona.json、materials/<mid>.json + 原文件。
"""
from __future__ import annotations

import asyncio
import contextlib
import io
import json
import logging
import re
import shutil
import subprocess
import tempfile
import time
import uuid
from pathlib import Path
from typing import Any

import httpx
from fastapi import APIRouter, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import JSONResponse, StreamingResponse

import cluster
import llm_router
import ot_pool
import personas as templates
import tts
from config import LLMS, current_provider

log = logging.getLogger("persona")
router = APIRouter(prefix="/v1/personas")
DIR = Path("/w/state/personas")
JOBS: dict[str, dict] = {}
INLINE_TOKENS = 3000
KINDS = ("doc", "image", "audio", "video", "url")
_BUSY: set[str] = set()

SCHEMA_HINT = """{
 "summary": "这份素材讲了什么（一两句）",
 "identity": {"name": "", "role": "", "org": "", "intro": "一句话自我介绍"},
 "background": ["学历/经历/成就等要点"],
 "style": {"tone": "", "pace": "", "catchphrases": [], "address_user": "", "examples": [{"user": "", "bot": ""}]},
 "knowledge": {"facts": ["可以直接回答的事实要点"], "faq": [{"q": "", "a": ""}], "products": [{"name": "", "points": []}]},
 "boundaries": {"forbidden": [], "fallback": "", "handoff": ""},
 "recommend": {"voice_hint": "推荐什么样的声音（性别/年龄/气质）", "avatar_hint": "着装、场景", "opening": "开场白"}
}"""


# ---------- 存储 ----------

def _pdir(pid: str) -> Path:
    return DIR / pid


def _load(pid: str) -> dict:
    p = _pdir(pid) / "persona.json"
    if not p.is_file():
        raise HTTPException(404, "persona not found")
    return json.loads(p.read_text("utf-8"))


def _save(p: dict) -> None:
    d = _pdir(p["id"])
    d.mkdir(parents=True, exist_ok=True)
    tmp = d / "persona.tmp"
    tmp.write_text(json.dumps(p, ensure_ascii=False, indent=1), "utf-8")
    tmp.replace(d / "persona.json")


def _materials(pid: str) -> list[dict]:
    md = _pdir(pid) / "materials"
    out = []
    for f in sorted(md.glob("*.json")) if md.is_dir() else []:
        with contextlib.suppress(OSError, ValueError):
            out.append(json.loads(f.read_text("utf-8")))
    return sorted(out, key=lambda m: m["ts"])


def _save_material(pid: str, m: dict) -> None:
    md = _pdir(pid) / "materials"
    md.mkdir(parents=True, exist_ok=True)
    tmp = md / f"{m['id']}.tmp"
    tmp.write_text(json.dumps(m, ensure_ascii=False, indent=1), "utf-8")
    tmp.replace(md / f"{m['id']}.json")


def list_all() -> list[dict]:
    out = []
    for d in sorted(DIR.glob("ps_*")) if DIR.is_dir() else []:
        with contextlib.suppress(OSError, ValueError, HTTPException):
            p = _load(d.name)
            mats = _materials(d.name)
            cur = _version(p, None)
            out.append({"persona_id": p["id"], "name": p["name"], "status": p["status"], "version": p["version"],
                        "biz_ref": p.get("biz_ref"), "materials": len(mats),
                        "materials_running": sum(1 for m in mats if m["status"] == "running"),
                        "knowledge_mode": (cur or {}).get("compiled", {}).get("knowledge_mode"),
                        "prompt_tokens": (cur or {}).get("compiled", {}).get("prompt_tokens"),
                        "updated": p.get("updated"), "building": p["id"] in _BUSY})
    return out


def _version(p: dict, v: int | None) -> dict | None:
    vs = p.get("versions") or []
    if not vs:
        return None
    if v is None:
        return vs[-1]
    return next((x for x in vs if x["version"] == v), None)


def _job(kind: str, pid: str, **extra) -> dict:
    j = {"id": uuid.uuid4().hex[:12], "kind": kind, "persona_id": pid, "state": "running", "step": "排队中", "error": None,
         "result": None, "t0": time.time(), **extra}
    JOBS[j["id"]] = j
    for k in [k for k, x in JOBS.items() if time.time() - x["t0"] > 3600 and x["state"] != "running"]:
        JOBS.pop(k, None)
    return j


def _ntok(s: str) -> int:
    try:
        import tiktoken
        return len(tiktoken.get_encoding("cl100k_base").encode(s))
    except Exception:  # noqa: BLE001
        return len(s) // 2


def _json_from(text: str) -> dict:
    """模型回的 JSON 可能裹在 ```json 里或带解释，尽量抠出来。"""
    text = text.strip()
    m = re.search(r"```(?:json)?\s*(\{.*\})\s*```", text, re.S)
    if m:
        text = m.group(1)
    else:
        i, j = text.find("{"), text.rfind("}")
        if i >= 0 and j > i:
            text = text[i:j + 1]
    return json.loads(text)


def _provider() -> str:
    p = current_provider()
    return p if LLMS[p].ready else next(k for k, c in LLMS.items() if c.ready)


# ---------- 素材 → 文本 ----------

def _text_from_doc(name: str, raw: bytes) -> str:
    ext = Path(name).suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        return "\n".join((pg.extract_text() or "") for pg in PdfReader(io.BytesIO(raw)).pages)
    if ext == ".docx":
        import docx
        d = docx.Document(io.BytesIO(raw))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if ext in (".html", ".htm"):
        from bs4 import BeautifulSoup
        return BeautifulSoup(raw, "html.parser").get_text("\n")
    for enc in ("utf-8", "gb18030", "utf-16"):
        with contextlib.suppress(UnicodeDecodeError):
            return raw.decode(enc)
    return raw.decode("utf-8", errors="ignore")


async def _text_from_url(url: str) -> str:
    from bs4 import BeautifulSoup
    async with httpx.AsyncClient(timeout=30, follow_redirects=True, trust_env=False,
                                 headers={"User-Agent": "Mozilla/5.0 talk-persona"}) as c:
        r = await c.get(url)
        r.raise_for_status()
    soup = BeautifulSoup(r.text, "html.parser")
    for t in soup(["script", "style", "nav", "footer", "header"]):
        t.decompose()
    return re.sub(r"\n{3,}", "\n\n", soup.get_text("\n")).strip()


async def _describe_images(images: list[bytes], hint: str) -> str:
    prompt = ("逐张描述这些图片：里面的文字（原样抄录）、人物（性别、年龄段、着装、表情）、产品/场景/背景、整体风格。"
              "用中文，分条，不要寒暄。" + (f"\n上传者备注：{hint}" if hint else ""))
    return await llm_router.complete(_provider(), [{"role": "user", "content": prompt}], images=images, max_tokens=1500)


def _ffmpeg(args: list[str], data: bytes) -> bytes:
    p = subprocess.run(["ffmpeg", "-v", "error", "-i", "pipe:0", *args], input=data, capture_output=True, timeout=600)
    if p.returncode:
        raise RuntimeError(p.stderr.decode(errors="ignore")[-200:])
    return p.stdout


async def _transcribe(pcm16: bytes) -> str:
    """分 50 秒一段送节点识别（控制面没有识别模型）。"""
    node = ot_pool.first().node
    step = 16000 * 2 * 50
    parts = []
    async with httpx.AsyncClient(headers=cluster.HDR, timeout=120, trust_env=False) as c:
        for i in range(0, len(pcm16), step):
            r = await c.post(f"{node.url}/internal/transcribe", content=pcm16[i:i + step])
            r.raise_for_status()
            if t := r.json().get("text"):
                parts.append(t)
    return "\n".join(parts)


async def _text_from_av(raw: bytes, video: bool, hint: str, job: dict) -> str:
    job["step"] = "转码音频"
    pcm = await asyncio.to_thread(_ffmpeg, ["-vn", "-f", "s16le", "-ac", "1", "-ar", "16000", "pipe:1"], raw)
    if len(pcm) > 16000 * 2 * 3600:
        raise ValueError("音视频超过 60 分钟")
    job["step"] = f"语音识别（{len(pcm) / 32000 / 60:.0f} 分钟）"
    text = await _transcribe(pcm)
    if video:
        job["step"] = "抽关键帧看着装场景"
        with contextlib.suppress(Exception):
            frames = []
            with tempfile.TemporaryDirectory() as d:
                src = Path(d, "in.bin")
                src.write_bytes(raw)
                dur = float(subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0",
                                            str(src)], capture_output=True, text=True).stdout.strip() or 0)
                for k in (0.1, 0.5, 0.9):
                    out = subprocess.run(["ffmpeg", "-v", "error", "-ss", f"{dur * k:.2f}", "-i", str(src), "-frames:v", "1",
                                          "-vf", "scale=768:-2", "-f", "image2", "-c:v", "mjpeg", "pipe:1"],
                                         capture_output=True, timeout=60).stdout
                    if out:
                        frames.append(out)
            if frames:
                text += "\n\n[画面]\n" + await _describe_images(frames, hint)
    return text


# ---------- LLM 抽取 / 合并 / 编译 ----------

async def _extract(text: str, kind: str, hint: str) -> dict:
    text = text.strip()
    if not text:
        raise ValueError("没抽到任何文字")
    chunks = [text[i:i + 12000] for i in range(0, min(len(text), 60000), 12000)]
    results = []
    for n, ch in enumerate(chunks):
        sys_p = ("你在为一个数字人整理人设素材。从素材里抽取信息，严格按给定 JSON 结构输出，没有的字段留空/空数组，"
                 "不要编造；事实要点写成一句一条、可以直接拿来回答用户的话；style 只在素材能体现说话风格（如本人讲话的转写）时填。"
                 "只输出 JSON。")
        user = f"素材类型：{kind}" + (f"\n上传者备注：{hint}" if hint else "") + \
            (f"\n（第 {n + 1}/{len(chunks)} 段）" if len(chunks) > 1 else "") + \
            f"\n\nJSON 结构：\n{SCHEMA_HINT}\n\n素材内容：\n{ch}"
        out = await llm_router.complete(_provider(), [{"role": "system", "content": sys_p}, {"role": "user", "content": user}],
                                        max_tokens=3000)
        results.append(_json_from(out))
    if len(results) == 1:
        return results[0]
    merged = results[0]
    for r in results[1:]:          # 多段：知识拼起来，其余以第一段为准
        for k in ("facts", "faq", "products"):
            merged.setdefault("knowledge", {}).setdefault(k, []).extend((r.get("knowledge") or {}).get(k) or [])
        merged.setdefault("background", []).extend(r.get("background") or [])
    return merged


async def _merge(p: dict, mats: list[dict]) -> dict:
    manual = p.get("manual") or {}
    extracts = [{"material": m["name"], "kind": m["kind"], "extract": m.get("extract")} for m in mats
                if m["status"] == "done" and m.get("extract")]
    if not extracts and not manual:
        raise ValueError("没有可用的素材，也没有手填内容")
    sys_p = ("把多份素材抽取结果和人工填写的内容合并成一个数字人的完整人设，严格按 JSON 结构输出。"
             "人工填写的优先级最高；冲突时以更具体、更新的为准；知识去重合并、同一事实只留一条；"
             "style.examples 给 3 条符合该人设口吻的问答示例（一问一答各一句）；recommend.opening 写一句开场白。"
             "不要编造素材里没有的事实。只输出 JSON。")
    user = f"JSON 结构：\n{SCHEMA_HINT}\n\n人工填写：\n{json.dumps(manual, ensure_ascii=False)}\n\n素材抽取结果：\n" + \
        json.dumps(extracts, ensure_ascii=False)[:60000]
    out = await llm_router.complete(_provider(), [{"role": "system", "content": sys_p}, {"role": "user", "content": user}],
                                    max_tokens=4000)
    data = _json_from(out)
    data.pop("summary", None)
    return data


def _pick_voice(hint: str) -> str:
    """recommend.voice_hint（"三十岁左右温柔女声"）→ 一个本地预设音色 ID。"""
    h = hint or ""
    female = any(k in h for k in ("女", "姐", "妈", "阿姨", "奶奶"))
    if any(k in h for k in ("童", "孩")):
        return "vox_girl_kid" if female else "vox_boy_kid"
    if any(k in h for k in ("少年", "少女", "学生")):
        return "vox_teen_girl" if female else "vox_m_sunny"
    if any(k in h for k in ("老", "奶奶", "爷爷", "长者")):
        return "vox_elder_f" if female else "vox_elder_m"   # 真实老人音色（vox_grandpa 是京腔逗乐款）
    if any(k in h for k in ("沉稳", "专业", "高管", "成熟", "中年")):
        return "vox_f_anchor" if female else "vox_m_mature"
    if any(k in h for k in ("播", "主持")):
        return "vox_f_anchor" if female else "vox_m_radio"
    return "vox_f_gentle" if female or not any(k in h for k in ("男", "哥", "叔", "先生")) else "vox_m_gentle"


def compile_prompt(data: dict) -> dict:
    idt, st, kn, bd, rc = (data.get(k) or {} for k in ("identity", "style", "knowledge", "boundaries", "recommend"))
    lines = []
    who = "、".join(x for x in (idt.get("name"), idt.get("role"), idt.get("org")) if x)
    if who:
        lines.append(f"你是{who}。" + (idt.get("intro") or ""))
    elif idt.get("intro"):
        lines.append(idt["intro"])
    if bg := data.get("background"):
        lines.append("背景：" + "；".join(str(x) for x in bg[:8]) + "。")
    sty = []
    if st.get("tone"):
        sty.append(f"语气{st['tone']}")
    if st.get("pace"):
        sty.append(st["pace"])
    if st.get("address_user"):
        sty.append(f"称呼对方为“{st['address_user']}”")
    if st.get("catchphrases"):
        sty.append("口头禅有：" + "、".join(f"“{c}”" for c in st["catchphrases"][:5]))
    if sty:
        lines.append("说话风格：" + "，".join(sty) + "。")
    if ex := st.get("examples"):
        lines.append("示例：\n" + "\n".join(f"用户：{e.get('user', '')}\n你：{e.get('bot', '')}" for e in ex[:3] if isinstance(e, dict)))
    if bd.get("forbidden"):
        lines.append("不要谈及：" + "；".join(str(x) for x in bd["forbidden"]) + "。")
    if bd.get("fallback"):
        lines.append(f"超出你知道的范围时：{bd['fallback']}")
    if bd.get("handoff"):
        lines.append(f"需要转人工时说：{bd['handoff']}")
    core = "\n".join(lines) + "\n" + templates.RULES
    know = []
    for f in kn.get("facts") or []:
        know.append(f"- {f}")
    for q in kn.get("faq") or []:
        if isinstance(q, dict) and q.get("q"):
            know.append(f"- 问：{q['q']} 答：{q.get('a', '')}")
    for pr in kn.get("products") or []:
        if isinstance(pr, dict) and pr.get("name"):
            know.append(f"- {pr['name']}：" + "；".join(str(x) for x in pr.get("points") or []))
    full = core + ("\n\n你掌握的信息（据此回答，不要编造）：\n" + "\n".join(know) if know else "")
    n_full, n_core = _ntok(full), _ntok(core)
    if n_full <= INLINE_TOKENS:
        return {"system_prompt": full, "prompt_tokens": n_full, "knowledge_mode": "inline", "knowledge_items": len(know)}
    return {"system_prompt": core + "\n\n（本会话的知识由外部检索提供，回答时以检索结果为准。）", "prompt_tokens": n_core,
            "knowledge_mode": "external", "knowledge_items": len(know), "knowledge_tokens": n_full - n_core}


def _new_version(p: dict, data: dict, note: str) -> dict:
    compiled = compile_prompt(data)
    rc = data.get("recommend") or {}
    if not rc.get("voice_id"):
        rc["voice_id"] = _pick_voice(rc.get("voice_hint", ""))
        data["recommend"] = rc
    p["version"] = int(p.get("version") or 0) + 1
    ver = {"version": p["version"], "ts": time.time(), "note": note, "data": data, "compiled": compiled}
    p.setdefault("versions", []).append(ver)
    p["versions"] = p["versions"][-20:]
    p["status"] = "draft"
    p["updated"] = time.time()
    _save(p)
    return ver


def compiled_for(pid: str, version: int | None = None) -> tuple[dict, dict]:
    """给会话用：(persona, 版本)。"""
    p = _load(pid)
    v = _version(p, version)
    if not v:
        raise HTTPException(400, "人设还没有 build 过")
    return p, v


# ---------- 后台任务 ----------

async def _run_material(pid: str, m: dict, raw: bytes | None, job: dict) -> None:
    hint = m.get("hint") or ""
    try:
        job["step"] = "读取素材"
        if m["kind"] == "url":
            text = await _text_from_url(m["source"])
        elif m["kind"] == "doc":
            text = await asyncio.to_thread(_text_from_doc, m["name"], raw or b"")
        elif m["kind"] == "image":
            job["step"] = "看图"
            text = await _describe_images([raw or b""], hint)
        else:
            text = await _text_from_av(raw or b"", m["kind"] == "video", hint, job)
        m["chars"] = len(text)
        (_pdir(pid) / "materials" / f"{m['id']}.txt").write_text(text[:200000], "utf-8")
        job["step"] = "抽取要点"
        ex = await _extract(text, m["kind"], hint)
        m.update(status="done", extract=ex, summary=str(ex.get("summary") or "")[:200], error=None)
        job.update(state="done", step="完成", result={"material_id": m["id"], "summary": m["summary"], "chars": m["chars"]})
        cluster.note("persona", f"人设 {pid} 素材「{m['name']}」抽取完成（{m['chars']} 字）")
    except Exception as e:  # noqa: BLE001
        log.exception("material %s failed", m["id"])
        m.update(status="failed", error=str(e)[:300])
        job.update(state="error", step="失败", error=str(e)[:300])
        cluster.note("persona", f"人设 {pid} 素材「{m['name']}」抽取失败：{str(e)[:60]}")
    finally:
        _save_material(pid, m)
        import public_api
        p = _load(pid)
        ctx = {"callback_url": m.get("callback_url") or p.get("callback_url"), "biz_ref": p.get("biz_ref")}
        ev = "material.extracted" if m["status"] == "done" else "material.failed"
        public_api.emit(ctx, ev, {"persona_id": pid, "material_id": m["id"], "summary": m.get("summary"), "error": m.get("error")})


async def _run_build(pid: str, job: dict) -> None:
    _BUSY.add(pid)
    try:
        p = _load(pid)
        job["step"] = "合并素材"
        data = await _merge(p, _materials(pid))
        ver = _new_version(p, data, "build")
        job.update(state="done", step="完成", result={"version": ver["version"], "compiled": ver["compiled"]})
        cluster.note("persona", f"人设「{p['name']}」build 完成 → v{ver['version']}（{ver['compiled']['knowledge_mode']}，{ver['compiled']['prompt_tokens']} tokens）")
        import public_api
        public_api.emit({"callback_url": p.get("callback_url"), "biz_ref": p.get("biz_ref")}, "persona.built",
                        {"persona_id": pid, "version": ver["version"], "knowledge_mode": ver["compiled"]["knowledge_mode"],
                         "prompt_tokens": ver["compiled"]["prompt_tokens"]})
    except Exception as e:  # noqa: BLE001
        log.exception("build %s failed", pid)
        job.update(state="error", step="失败", error=str(e)[:300])
        cluster.note("persona", f"人设 {pid} build 失败：{str(e)[:60]}")
    finally:
        _BUSY.discard(pid)


def startup() -> None:
    """控制面重启：抽到一半的素材标失败（原文件还在，可重传或 retry）。"""
    for d in DIR.glob("ps_*") if DIR.is_dir() else []:
        for m in _materials(d.name):
            if m["status"] == "running":
                m.update(status="failed", error="控制面重启中断，请重新上传或 retry")
                _save_material(d.name, m)


# ---------- 演示页一键流水线：建人设 + 传素材 + 等抽取 + build ----------

PIPELINES: dict[str, dict] = {}


def start_pipeline(name: str, hint: str, files: list[tuple[str, str, bytes]], urls: list[str], seed: dict | None = None) -> dict:
    """files: [(kind, 文件名, 字节)]。返回流水线任务（演示页轮询用）。"""
    pid = f"ps_{uuid.uuid4().hex[:10]}"
    p = {"id": pid, "name": name, "biz_ref": "demo", "callback_url": None, "status": "draft", "version": 0, "versions": [],
         "manual": seed or {}, "created": time.time(), "updated": time.time()}
    _save(p)
    cluster.note("persona", f"演示页新建人设「{name}」{pid}（{len(files)} 个文件，{len(urls)} 个网址）")
    pl = {"id": uuid.uuid4().hex[:12], "persona_id": pid, "name": name, "state": "running", "step": "抽取素材",
          "materials": [], "t0": time.time(), "error": None, "version": None, "compiled": None}
    PIPELINES[pl["id"]] = pl
    for k in [k for k, x in PIPELINES.items() if time.time() - x["t0"] > 3600]:
        PIPELINES.pop(k, None)
    cluster.spawn(_run_pipeline(pl, hint, files, urls))
    return pl


async def _run_pipeline(pl: dict, hint: str, files: list[tuple[str, str, bytes]], urls: list[str]) -> None:
    pid = pl["persona_id"]
    jobs: list[tuple[dict, dict]] = []
    (_pdir(pid) / "materials").mkdir(parents=True, exist_ok=True)
    for kind, name, raw in files:
        mid = f"mat_{uuid.uuid4().hex[:10]}"
        m = {"id": mid, "kind": kind, "name": name, "source": None, "hint": hint[:200], "status": "running", "size": len(raw),
             "ts": time.time()}
        (_pdir(pid) / "materials" / f"{mid}{Path(name).suffix.lower()[:8] or '.bin'}").write_bytes(raw)
        _save_material(pid, m)
        job = _job("material", pid, material_id=mid)
        jobs.append((m, job))
        cluster.spawn(_run_material(pid, m, raw, job))
    for u in urls:
        mid = f"mat_{uuid.uuid4().hex[:10]}"
        m = {"id": mid, "kind": "url", "name": u, "source": u, "hint": hint[:200], "status": "running", "size": 0, "ts": time.time()}
        _save_material(pid, m)
        job = _job("material", pid, material_id=mid)
        jobs.append((m, job))
        cluster.spawn(_run_material(pid, m, None, job))
    try:
        while True:
            pl["materials"] = [{"name": m["name"], "kind": m["kind"], "status": "running" if j["state"] == "running" else
                                ("done" if j["state"] == "done" else "failed"), "step": j["step"],
                                "summary": (j.get("result") or {}).get("summary") or j.get("error") or ""} for m, j in jobs]
            if all(j["state"] != "running" for _, j in jobs):
                break
            await asyncio.sleep(2)
        if jobs and not any(j["state"] == "done" for _, j in jobs) and not (_load(pid).get("manual")):
            raise RuntimeError("素材全部抽取失败：" + "；".join(str(j.get("error"))[:60] for _, j in jobs))
        pl["step"] = "合并编译人设"
        bj = _job("build", pid)
        await _run_build(pid, bj)
        if bj["state"] != "done":
            raise RuntimeError(bj.get("error") or "build 失败")
        pl.update(state="done", step="完成", version=bj["result"]["version"], compiled=bj["result"]["compiled"])
    except Exception as e:  # noqa: BLE001
        pl.update(state="error", step="失败", error=str(e)[:300])


# ---------- 接口 ----------

def _key(request: Request) -> None:
    import public_api
    public_api.require_key(request)


def _public(p: dict, version: int | None = None, full: bool = True) -> dict:
    v = _version(p, version)
    out = {"persona_id": p["id"], "name": p["name"], "status": p["status"], "version": p["version"], "biz_ref": p.get("biz_ref"),
           "created": p.get("created"), "updated": p.get("updated"), "manual": p.get("manual") or {},
           "versions": [{"version": x["version"], "ts": x["ts"], "note": x["note"]} for x in p.get("versions") or []]}
    if v and full:
        out |= {"data": v["data"], "compiled": v["compiled"], "this_version": v["version"]}
    return out


@router.post("", status_code=201)
async def create(request: Request, body: dict) -> dict:
    _key(request)
    name = str(body.get("name") or "").strip()[:40]
    if not name:
        raise HTTPException(400, "name 必填")
    pid = f"ps_{uuid.uuid4().hex[:10]}"
    p = {"id": pid, "name": name, "biz_ref": body.get("biz_ref"), "callback_url": body.get("callback_url"), "status": "draft",
         "version": 0, "versions": [], "manual": body.get("seed") or {}, "created": time.time(), "updated": time.time()}
    _save(p)
    if p["manual"]:
        _new_version(p, dict(p["manual"]), "seed")
    cluster.note("persona", f"新建人设「{name}」{pid}")
    return _public(p)


@router.get("")
async def index(request: Request) -> list[dict]:
    _key(request)
    return list_all()


@router.get("/{pid}")
async def get(pid: str, request: Request, version: int | None = None) -> dict:
    _key(request)
    return _public(_load(pid), version)


@router.patch("/{pid}")
async def patch(pid: str, request: Request, body: dict) -> dict:
    """人工改字段：和当前版本的数据合并 → 新版本重新编译。body 直接给 schema 里的字段（identity/style/knowledge/…）或 name。"""
    _key(request)
    p = _load(pid)
    if "name" in body:
        p["name"] = str(body.pop("name"))[:40]
    if "callback_url" in body:
        p["callback_url"] = body.pop("callback_url")
    fields = {k: v for k, v in body.items() if k in ("identity", "background", "style", "knowledge", "boundaries", "recommend")}
    if fields:
        manual = p.get("manual") or {}
        for k, v in fields.items():
            manual[k] = {**manual.get(k, {}), **v} if isinstance(v, dict) and isinstance(manual.get(k), dict) else v
        p["manual"] = manual
        cur = (_version(p, None) or {}).get("data") or {}
        data = json.loads(json.dumps(cur))
        for k, v in fields.items():
            data[k] = {**data.get(k, {}), **v} if isinstance(v, dict) and isinstance(data.get(k), dict) else v
        _new_version(p, data, "manual")
    else:
        _save(p)
    return _public(p)


@router.delete("/{pid}")
async def delete(pid: str, request: Request) -> dict:
    _key(request)
    _load(pid)
    shutil.rmtree(_pdir(pid), ignore_errors=True)
    return {"ok": True}


@router.post("/{pid}/materials", status_code=202)
async def add_material(pid: str, request: Request, kind: str = Form(...), file: UploadFile | None = File(None),
                       url: str = Form(""), hint: str = Form(""), callback_url: str = Form("")) -> dict:
    _key(request)
    p = _load(pid)
    if kind not in KINDS:
        raise HTTPException(400, f"kind 只能是 {'/'.join(KINDS)}")
    raw = None
    if kind == "url":
        if not url.startswith(("http://", "https://")):
            raise HTTPException(400, "url 不合法")
        name = url
    else:
        if file is None:
            raise HTTPException(400, "缺文件")
        raw = await file.read()
        if not raw or len(raw) > 200 * 1024 * 1024:
            raise HTTPException(400, "文件为空或超过 200MB")
        name = file.filename or kind
    mid = f"mat_{uuid.uuid4().hex[:10]}"
    m = {"id": mid, "kind": kind, "name": name, "source": url or None, "hint": hint[:200], "status": "running",
         "size": len(raw) if raw else 0, "ts": time.time(), "callback_url": callback_url or None}
    (_pdir(pid) / "materials").mkdir(parents=True, exist_ok=True)
    if raw:
        (_pdir(pid) / "materials" / f"{mid}{Path(name).suffix.lower()[:8] or '.bin'}").write_bytes(raw)
    _save_material(pid, m)
    job = _job("material", pid, material_id=mid)
    cluster.spawn(_run_material(pid, m, raw, job))
    return {"material_id": mid, "job_id": job["id"], "persona_id": pid}


@router.get("/{pid}/materials")
async def materials(pid: str, request: Request) -> list[dict]:
    _key(request)
    _load(pid)
    return [{k: m.get(k) for k in ("id", "kind", "name", "hint", "status", "summary", "error", "chars", "size", "ts")}
            for m in _materials(pid)]


@router.get("/{pid}/materials/{mid}")
async def material(pid: str, request: Request, mid: str) -> dict:
    _key(request)
    _load(pid)
    m = next((x for x in _materials(pid) if x["id"] == mid), None)
    if not m:
        raise HTTPException(404, "material not found")
    txt = _pdir(pid) / "materials" / f"{mid}.txt"
    return m | {"text_preview": txt.read_text("utf-8")[:2000] if txt.is_file() else ""}


@router.delete("/{pid}/materials/{mid}")
async def delete_material(pid: str, request: Request, mid: str) -> dict:
    _key(request)
    _load(pid)
    md = _pdir(pid) / "materials"
    n = 0
    for f in md.glob(f"{mid}*"):
        f.unlink(missing_ok=True)
        n += 1
    if not n:
        raise HTTPException(404, "material not found")
    return {"ok": True, "需要": "重新 build 才会从人设里去掉这份素材的内容"}


@router.post("/{pid}/materials/{mid}/retry", status_code=202)
async def retry_material(pid: str, request: Request, mid: str) -> dict:
    _key(request)
    _load(pid)
    m = next((x for x in _materials(pid) if x["id"] == mid), None)
    if not m:
        raise HTTPException(404, "material not found")
    raw = None
    if m["kind"] != "url":
        f = next((x for x in (_pdir(pid) / "materials").glob(f"{mid}.*") if x.suffix not in (".json", ".txt")), None)
        if not f:
            raise HTTPException(410, "原文件已不在，请重新上传")
        raw = f.read_bytes()
    m.update(status="running", error=None)
    _save_material(pid, m)
    job = _job("material", pid, material_id=mid)
    cluster.spawn(_run_material(pid, m, raw, job))
    return {"material_id": mid, "job_id": job["id"]}


@router.post("/{pid}/build", status_code=202)
async def build(pid: str, request: Request) -> dict:
    _key(request)
    _load(pid)
    if pid in _BUSY:
        raise HTTPException(409, "正在 build")
    if any(m["status"] == "running" for m in _materials(pid)):
        raise HTTPException(409, "还有素材在抽取，等它们完成再 build")
    job = _job("build", pid)
    cluster.spawn(_run_build(pid, job))
    return {"job_id": job["id"], "persona_id": pid}


@router.post("/{pid}/chat")
async def chat(pid: str, request: Request, body: dict):
    """文本试聊（不占 GPU）：用当前编译结果走 LLM，OpenAI 风格 SSE 流。"""
    _key(request)
    _, v = compiled_for(pid, body.get("version"))
    msgs = [{"role": "system", "content": v["compiled"]["system_prompt"]}] + \
        [m for m in (body.get("messages") or []) if m.get("role") in ("user", "assistant")][-20:]
    return StreamingResponse(llm_router.stream(_provider(), msgs, int(body.get("max_tokens") or 400)),
                             media_type="text/event-stream")


@router.post("/{pid}/approve")
async def approve(pid: str, request: Request, body: dict | None = None) -> dict:
    _key(request)
    p, v = compiled_for(pid, (body or {}).get("version"))
    p["status"] = "approved"
    p["approved_version"] = v["version"]
    p["updated"] = time.time()
    _save(p)
    cluster.note("persona", f"人设「{p['name']}」v{v['version']} 已审核通过")
    return {"persona_id": pid, "status": "approved", "version": v["version"], "compiled": v["compiled"]}


@router.get("/{pid}/export")
async def export(pid: str, request: Request, version: int | None = None) -> dict:
    """给后端建 RAG：编译后的 prompt + 全量知识条目。"""
    _key(request)
    p, v = compiled_for(pid, version)
    kn = v["data"].get("knowledge") or {}
    return {"persona_id": pid, "version": v["version"], "system_prompt": v["compiled"]["system_prompt"],
            "knowledge_mode": v["compiled"]["knowledge_mode"], "knowledge": kn, "recommend": v["data"].get("recommend") or {},
            "sources": [{"material_id": m["id"], "name": m["name"], "kind": m["kind"]} for m in _materials(pid) if m["status"] == "done"]}
